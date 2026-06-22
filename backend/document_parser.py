import csv
import os

import fitz


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx"}


def extract_text_from_pdf(file_path):
    document = fitz.open(file_path)
    pages = []
    try:
        for page_index, page in enumerate(document):
            text = page.get_text()
            if text.strip():
                pages.append({"page_number": page_index + 1, "text": text})
    finally:
        document.close()
    return pages


def extract_text_from_txt(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        text = file.read()
    return [{"page_number": 1, "text": text}]


def extract_text_from_csv(file_path):
    lines = []
    with open(file_path, "r", encoding="utf-8-sig", errors="ignore", newline="") as file:
        for row in csv.reader(file):
            cleaned = [cell.strip() for cell in row if cell and cell.strip()]
            if cleaned:
                lines.append(" | ".join(cleaned))
    return [{"page_number": 1, "text": "\n".join(lines)}]


def extract_text_from_docx(file_path):
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError(
            "DOCX parsing requires python-docx. Run: pip install python-docx"
        ) from error

    document = Document(file_path)
    blocks = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return [{"page_number": 1, "text": "\n".join(blocks)}]


def parse_document(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension}")

    if extension == ".pdf":
        pages = extract_text_from_pdf(file_path)
    elif extension in {".txt", ".md"}:
        pages = extract_text_from_txt(file_path)
    elif extension == ".csv":
        pages = extract_text_from_csv(file_path)
    else:
        pages = extract_text_from_docx(file_path)

    pages = [page for page in pages if page.get("text", "").strip()]
    if not pages:
        raise ValueError("The document does not contain extractable text")
    return pages

